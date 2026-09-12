# pages/app_limpieza_excel.py — v3.13.3
# Cambios sobre tu baseline v3.13.1:
# - GEOCODER_TIMEOUT bajado a 8 (antes 20).
# - reverse_best_multi ahora SOLO usa Nominatim (geopy + HTTP) + OpenCage.
#   (Se eliminaron del flujo _reverse_locationiq y _reverse_bigdatacloud).
# - FIX v3.13.3: geocoding principal ahora usa reverse_address() con caché.
# - FIX v3.13.3: pluscode_label usa caché administrativa por coordenada.
# - FIX v3.13.3: Direccion_final resuelve SOLO coordenadas únicas faltantes.
#
# El resto de la lógica (parsing de fechas, dedupe DATOS, estadísticas,
# plus_repo, geo_cache, Bloque 7, etc.) se mantiene IGUAL a tu baseline.

import os, io, re, time, sqlite3, unicodedata
from contextlib import closing
from functools import lru_cache
from typing import Tuple, Optional, Dict, List
from datetime import datetime, timedelta, date, time as dtime

import pandas as pd
import requests
import streamlit as st
from core.limpieza_offline import offline_address
from guardian import login_guard
from suite_nav import render_suite_sidebar  # ✅ Navegación unificada de la suite
from ui.styles import apply_theme, info_panel, page_header, section_title
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font

# ===========================
# CONFIG STREAMLIT
# ===========================
st.set_page_config(
    page_title="Go Mapper - Limpieza Automática",
    page_icon="🛰️",
    layout="wide",
    menu_items={"Get help": None, "Report a bug": None, "About": None}
)

# ===========================
# GUARDIÁN DE ACCESO (nuevo)
# ===========================
# Usa el guardian.py con [users] en secrets.toml
login_guard("Limpieza de CDR")

# ===========================
# UI BLINDADA + NAVEGACIÓN
# ===========================
apply_theme()

# Navegación lateral estándar de la suite
try:
    render_suite_sidebar()
except Exception:
    pass

# ===========================
# DEPENDENCIAS CRÍTICAS
# ===========================
try:
    from openlocationcode import openlocationcode as olc
except Exception:
    st.error("Falta 'openlocationcode'. Agrega 'openlocationcode>=1.0.0' a requirements.txt y redeploy.")
    st.stop()

try:
    from geopy.geocoders import Nominatim
    from geopy.extra.rate_limiter import RateLimiter
except Exception:
    st.error("Falta 'geopy'. Agrega 'geopy>=2.4.1' a requirements.txt y redeploy.")
    st.stop()

# ===========================
# CONFIG GEOCODER BASE
# ===========================
CONTACT_EMAIL = "jgonzaloromerolugo@gmail.com"
CONTACT_URL = os.getenv("CONTACT_URL", "")
USER_AGENT = (
    f"go-mapper/3.13.3 ({CONTACT_EMAIL})"
    if not CONTACT_URL
    else f"go-mapper/3.13.3 (+{CONTACT_URL}; {CONTACT_EMAIL})"
)

NOMINATIM_URL = os.getenv("NOMINATIM_URL", "https://nominatim.openstreetmap.org").rstrip("/")


def _host_from_url(url: str) -> str:
    return re.sub(r"^https?://", "", url, flags=re.I).rstrip("/")


NOMINATIM_HOST = _host_from_url(NOMINATIM_URL)

# 🔧 Ajuste de rendimiento: timeout más bajo y tope de geocodes
GEOCODER_TIMEOUT = 8  # antes 20, ahora en rango 5–10 s como acordamos
MAX_UNIQUE_GEOCODES = 1000
COORD_PRECISION_CACHE = 6
CACHE_DB_PATH = "geo_cache.sqlite"
GEOCODE_ENABLED = True

# ===========================
# PLUS REPO (SQLite persistente)
# ===========================
PLUS_REPO_DIR = os.path.join("data", "plus_repo")
PLUS_REPO_DB = os.path.join(PLUS_REPO_DIR, "plus_repo.sqlite")
os.makedirs(PLUS_REPO_DIR, exist_ok=True)


def _pr_conn():
    con = sqlite3.connect(PLUS_REPO_DB)
    con.execute("PRAGMA journal_mode=WAL;")
    con.execute("PRAGMA synchronous=NORMAL;")
    return con


def init_plus_repo():
    with closing(_pr_conn()) as con:
        con.execute(
            """CREATE TABLE IF NOT EXISTS plus_codes(
            code TEXT PRIMARY KEY,
            lat REAL, lon REAL,
            code_len INT,
            updated_at INTEGER
        )"""
        )
        con.execute(
            """CREATE TABLE IF NOT EXISTS plus_names(
            code TEXT PRIMARY KEY,
            nombre TEXT,
            fuente TEXT,
            conf REAL,
            updated_at INTEGER
        )"""
        )
        con.execute(
            """CREATE TABLE IF NOT EXISTS near_names(
            lat5 INT, lon5 INT,
            nombre TEXT, fuente TEXT, conf REAL,
            UNIQUE(lat5, lon5)
        )"""
        )
        con.execute(
            """CREATE INDEX IF NOT EXISTS idx_near_names_latlon
                       ON near_names(lat5, lon5)"""
        )
        con.commit()


def _ll_key(lat: float, lon: float, places: int = 5) -> Tuple[int, int]:
    m = 10**places
    return (int(round(float(lat) * m)), int(round(float(lon) * m)))


def _decode_plus(code: str) -> Tuple[float, float, int]:
    area = olc.decode(code)
    lat = (area.latitudeLo + area.latitudeHi) / 2.0
    lon = (area.longitudeLo + area.longitudeHi) / 2.0
    return lat, lon, area.codeLength


def pr_get_nombre(code: str, lat: Optional[float] = None, lon: Optional[float] = None) -> Optional[Tuple[str, str, float]]:
    with closing(_pr_conn()) as con:
        r = con.execute("SELECT nombre, fuente, conf FROM plus_names WHERE code=?", (code,)).fetchone()
        if r:
            return r
        if (lat is not None) and (lon is not None):
            lat5, lon5 = _ll_key(lat, lon, 5)
            r = con.execute(
                "SELECT nombre, fuente, conf FROM near_names WHERE lat5=? AND lon5=?",
                (lat5, lon5),
            ).fetchone()
            if r:
                return r
    return None


def pr_save_plus(code: str, lat: float, lon: float, code_len: int):
    ts = int(time.time())
    with closing(_pr_conn()) as con:
        con.execute(
            """INSERT OR REPLACE INTO plus_codes(code,lat,lon,code_len,updated_at)
                       VALUES (?,?,?,?,?)""",
            (code, lat, lon, code_len, ts),
        )
        con.commit()


def pr_save_nombre(
    code: str,
    nombre: str,
    fuente: str = "cache",
    conf: float = 0.9,
    lat: Optional[float] = None,
    lon: Optional[float] = None,
):
    ts = int(time.time())
    with closing(_pr_conn()) as con:
        con.execute(
            """INSERT OR REPLACE INTO plus_names(code,nombre,fuente,conf,updated_at)
                       VALUES (?,?,?,?,?)""",
            (code, nombre, fuente, conf, ts),
        )
        if (lat is not None) and (lon is not None):
            lat5, lon5 = _ll_key(lat, lon, 5)
            con.execute(
                """INSERT OR REPLACE INTO near_names(lat5,lon5,nombre,fuente,conf)
                           VALUES (?,?,?,?,?)""",
                (lat5, lon5, nombre, fuente, conf),
            )
        con.commit()


def pr_get_near_name(lat: float, lon: float) -> str:
    with closing(_pr_conn()) as con:
        lat5, lon5 = _ll_key(lat, lon, 5)
        r = con.execute(
            "SELECT nombre FROM near_names WHERE lat5=? AND lon5=?",
            (lat5, lon5),
        ).fetchone()
        return r[0] if r and r[0] else ""


# ===========================
# UTILIDADES
# ===========================
def progress_section(progress, pct, msg):
    progress.progress(min(max(int(pct), 0), 100), text=msg)


def norm(s: str) -> str:
    if s is None:
        return ""
    s = str(s).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return s.lower()


def _sanitize_filename(name: str, ext: str = ".xlsx") -> str:
    name = re.sub(r'[\\/:*?"<>|]+', " ", str(name)).strip()
    name = re.sub(r"\s+", " ", name)
    if not name.lower().endswith(ext.lower()):
        name = f"{name}{ext}"
    return name


def _clean_number(x: str) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    return re.sub(r"\D+", "", str(x))


# ===========================
# DMS / ARCHIVOS / AZIMUTH
# ===========================
def dms_to_decimal(value):
    if pd.isna(value):
        return None
    value = str(value).strip()
    if re.search(r"[°'\"NSEOWeo]", value, re.I):
        match = re.findall(
            r"(\d+)[°\s]+(\d+)?['\s]*([\d\.]+)?\"?\s*([NSEOWO])?",
            value,
            re.I,
        )
        if not match:
            return None
        deg, minu, sec, hemi = match[0]
        deg = float(deg)
        minu = float(minu or 0)
        sec = float(sec or 0)
        dec = deg + minu / 60 + sec / 3600
        if hemi and hemi.upper() in ["S", "W", "O"]:
            dec = -dec
        return dec
    try:
        return float(value.replace(",", "."))
    except:  # noqa: E722
        return None


def leer_archivo(file):
    name = (getattr(file, "name", "") or "").lower()
    if name.endswith(".csv"):
        return pd.read_csv(file, header=None, engine="python", sep=None, encoding_errors="replace")
    elif name.endswith(".xls"):
        return pd.read_excel(file, header=None, engine="xlrd")
    else:
        return pd.read_excel(file, header=None, engine="openpyxl")


# Azimuth robusto
AZI_CARDINALS = re.compile(r"[NSEOW]", re.I)  # O = Oeste
AZI_NUMBER = re.compile(r"-?\d+(?:\.\d+)?")

def parse_azimuth(val):
    """Acepta '240°', '240 º', '260deg', '260 O', '-', None… y regresa 0..360 (float)"""
    if pd.isna(val):
        return None
    s = str(val).strip().upper()
    if s in {"", "-", "NA", "N/A", "NULL"}:
        return None
    s = s.replace("º", "").replace("°", "").replace("DEG", "")
    s = AZI_CARDINALS.sub("", s)
    m = AZI_NUMBER.search(s)
    if not m:
        return None
    try:
        x = float(m.group(0))
        return round(x % 360, 3)
    except Exception:
        return None

# ===========================
# EXTRAER COORDENADAS
# ===========================
def try_split_two_coords(cell: str) -> Tuple[Optional[float], Optional[float]]:
    if cell is None or (isinstance(cell, float) and pd.isna(cell)):
        return (None, None)
    s = str(cell).strip()
    if not s:
        return (None, None)
    s_std = s.replace(";", ",").replace("/", ",")
    if "," in s_std:
        parts = [p.strip() for p in s_std.split(",", 1)]
        if len(parts) == 2:
            lat = dms_to_decimal(parts[0])
            lon = dms_to_decimal(parts[1])
            if (lat is not None) and (lon is not None):
                return (lat, lon)
    nums = re.findall(r"(-?\d+(?:[.,]\d+)?)", s)
    if len(nums) >= 2:
        lat = dms_to_decimal(nums[0])
        lon = dms_to_decimal(nums[1])
        if (lat is not None) and (lon is not None):
            return (lat, lon)
    blocks = re.findall(r"(\d+[^NSEWO]*[NSEWO])", s, flags=re.I)
    if len(blocks) >= 2:
        lat = dms_to_decimal(blocks[0])
        lon = dms_to_decimal(blocks[1])
        if (lat is not None) and (lon is not None):
            return (lat, lon)
    return (None, None)


def ensure_lat_lon_columns(df: pd.DataFrame) -> pd.DataFrame:
    if {"Latitud", "Longitud"}.issubset(df.columns):
        return df
    df2 = df.copy()
    name_patterns = [
        r"ubicaci[oó]n.*(latitud.*longitud|coordenadas)",
        r"coordenadas",
        r"geo.*(lat|long)",
        r"latitud\s*[/|-]\s*longitud",
        r"lat.*long",
    ]
    candidate_cols = []
    for col in df2.columns:
        c = str(col).strip().lower()
        if any(re.search(p, c, flags=re.I) for p in name_patterns):
            candidate_cols.append(col)
    sample_cols = candidate_cols if candidate_cols else list(df2.columns)
    found_col = None
    for col in sample_cols:
        series = df2[col].dropna().astype(str)
        if series.empty:
            continue
        ok = 0
        for v in series.head(200):
            lat, lon = try_split_two_coords(v)
            if (lat is not None) and (lon is not None):
                ok += 1
        if ok >= 10:
            found_col = col
            break
    if found_col is not None:
        lats, lons = [], []
        for v in df2[found_col].astype(str).tolist():
            lat, lon = try_split_two_coords(v)
            lats.append(lat)
            lons.append(lon)
        df2["Latitud"] = pd.Series(lats)
        df2["Longitud"] = pd.Series(lons)
        df2["Longitud"] = df2["Longitud"].apply(lambda x: -abs(x) if pd.notna(x) else x)
    return df2


def force_latlon_from_any(df: pd.DataFrame) -> pd.DataFrame:
    if "Latitud" not in df.columns:
        df["Latitud"] = None
    if "Longitud" not in df.columns:
        df["Longitud"] = None
    mask_valid = df["Latitud"].notna() & df["Longitud"].notna()
    if mask_valid.any():
        return df
    cand_cols = [c for c in df.columns if df[c].dtype == "object"]
    if not cand_cols:
        return df
    lat_list = list(df["Latitud"])
    lon_list = list(df["Longitud"])
    for i, row in df[cand_cols].astype(str).iterrows():
        if (lat_list[i] is None) or (lon_list[i] is None):
            joined = " | ".join([row[c] for c in cand_cols if isinstance(row[c], str)])
            lat, lon = try_split_two_coords(joined)
            if (lat is not None) and (lon is not None):
                lon = -abs(lon)
                lat_list[i], lon_list[i] = lat, lon
    df["Latitud"] = pd.Series(lat_list)
    df["Longitud"] = pd.Series(lon_list)
    return df

# ===========================
# GEOCODING: CADENA DE PROVEEDORES
# ===========================
ALT_URLS = [u.strip().rstrip("/") for u in os.getenv("NOMINATIM_ALT_URLS", "").split(",") if u.strip()]
ALT_HOSTS = [_host_from_url(u) for u in ALT_URLS]
HOSTS_CHAIN = [NOMINATIM_HOST] + ALT_HOSTS
URLS_CHAIN = [NOMINATIM_URL] + ALT_URLS

# Optional online credentials are resolved only when that provider is actually used.
OPENCAGE_KEY = os.getenv("OPENCAGE_API_KEY", "")
LOCATIONIQ_KEY = os.getenv("LOCATIONIQ_API_KEY", "")

def _optional_geocoder_key(name):
    try:
        return st.secrets.get(name, "")
    except (FileNotFoundError, KeyError):
        return ""


_GEOPY_REV_BY_HOST: Dict[str, RateLimiter] = {}

def _get_geopy_reverse_for_host(host: str) -> RateLimiter:
    if host not in _GEOPY_REV_BY_HOST:
        geolocator = Nominatim(user_agent=USER_AGENT, timeout=GEOCODER_TIMEOUT, domain=host)
        _GEOPY_REV_BY_HOST[host] = RateLimiter(
            geolocator.reverse,
            min_delay_seconds=1.8,
            max_retries=1,
            error_wait_seconds=2.0,
            swallow_exceptions=False,
        )
    return _GEOPY_REV_BY_HOST[host]


def _compose_from_addressdict(addr: dict) -> str:
    order = [
        "road",
        "neighbourhood",
        "suburb",
        "city",
        "town",
        "village",
        "municipality",
        "county",
        "state",
        "country",
    ]
    parts = []
    for k in order:
        v = addr.get(k)
        if v and v not in parts:
            parts.append(v)
    return ", ".join(parts)


def _reverse_http_base(base_url: str, lat, lon, lang="es", zoom=18, timeout=GEOCODER_TIMEOUT + 5) -> str:
    url = f"{base_url}/reverse"
    params = {
        "format": "jsonv2",
        "lat": f"{float(lat):.6f}",
        "lon": f"{float(lon):.6f}",
        "accept-language": lang,
        "zoom": zoom,
        "addressdetails": 1,
        "email": CONTACT_EMAIL,
    }
    headers = {"User-Agent": USER_AGENT}
    r = requests.get(url, params=params, headers=headers, timeout=timeout)
    r.raise_for_status()
    data = r.json()
    if isinstance(data, dict):
        if data.get("display_name"):
            return data["display_name"]
        if isinstance(data.get("address"), dict):
            return _compose_from_addressdict(data["address"])
    return ""


def _reverse_locationiq(lat, lon, lang="es"):
    api_key = LOCATIONIQ_KEY or _optional_geocoder_key("LOCATIONIQ_API_KEY")
    if not api_key:
        return ""
    url = "https://us1.locationiq.com/v1/reverse"
    params = {
        "key": api_key,
        "lat": f"{float(lat):.6f}",
        "lon": f"{float(lon):.6f}",
        "format": "json",
        "zoom": 18,
        "normalizeaddress": 1,
        "accept-language": lang,
    }
    r = requests.get(url, params=params, timeout=GEOCODER_TIMEOUT)
    r.raise_for_status()
    js = r.json()
    if isinstance(js, dict):
        if js.get("display_name"):
            return js["display_name"]
        if isinstance(js.get("address"), dict):
            return _compose_from_addressdict(js["address"])
    return ""


def _reverse_opencage(lat, lon, lang="es"):
    api_key = OPENCAGE_KEY or _optional_geocoder_key("OPENCAGE_API_KEY")
    if not api_key:
        return ""
    url = "https://api.opencagedata.com/geocode/v1/json"
    params = {
        "q": f"{float(lat):.6f},{float(lon):.6f}",
        "key": api_key,
        "language": lang,
        "no_annotations": 1,
        "limit": 1,
    }
    r = requests.get(url, params=params, timeout=GEOCODER_TIMEOUT)
    r.raise_for_status()
    js = r.json()
    if isinstance(js, dict) and js.get("results"):
        res = js["results"][0]
        if res.get("formatted"):
            return res["formatted"]
        if isinstance(res.get("components"), dict):
            return _compose_from_addressdict(res["components"])
    return ""


def _reverse_bigdatacloud(lat, lon, lang="es"):
    url = "https://api.bigdatacloud.net/data/reverse-geocode-client"
    params = {
        "latitude": float(lat),
        "longitude": float(lon),
        "localityLanguage": "es" if str(lang).lower().startswith("es") else "en",
    }
    r = requests.get(url, params=params, timeout=GEOCODER_TIMEOUT)
    r.raise_for_status()
    js = r.json()
    if not isinstance(js, dict):
        return ""
    parts = [
        js.get("locality") or js.get("city"),
        js.get("principalSubdivision"),
        js.get("countryName"),
    ]
    parts2 = [p or "" for p in parts]
    out = ", ".join([p for p in parts2 if p])
    return out

# --- NUEVO: helpers para PLUS corto + etiqueta administrativa ----------------
def _reverse_http_addr(base_url: str, lat, lon, lang="es", zoom=14, timeout=GEOCODER_TIMEOUT + 5) -> dict:
    url = f"{base_url}/reverse"
    params = {
        "format": "jsonv2",
        "lat": f"{float(lat):.6f}",
        "lon": f"{float(lon):.6f}",
        "accept-language": lang,
        "zoom": zoom,
        "addressdetails": 1,
        "email": CONTACT_EMAIL,
    }
    headers = {"User-Agent": USER_AGENT}
    r = requests.get(url, params=params, headers=headers, timeout=timeout)
    r.raise_for_status()
    data = r.json()
    return data.get("address", {}) if isinstance(data, dict) else {}


def _admin_label_from_addr(addr: dict) -> str:
    if not isinstance(addr, dict):
        return ""
    city = (
        addr.get("city")
        or addr.get("town")
        or addr.get("village")
        or addr.get("municipality")
        or addr.get("county")
    )
    state = addr.get("state") or addr.get("region") or addr.get("state_district")
    parts = [p for p in [city, state] if p]
    if parts:
        return ", ".join(parts)
    country = addr.get("country")
    parts2 = [p for p in [state, country] if p]
    return ", ".join(parts2)


def best_admin_label(lat, lon, lang="es") -> str:
    """Etiqueta administrativa ligera usando SOLO Nominatim HTTP (sin BigDataCloud)."""
    for Z in [14, 12, 10]:
        for base in URLS_CHAIN:
            try:
                addr = _reverse_http_addr(base, lat, lon, lang=lang, zoom=Z)
                lab = _admin_label_from_addr(addr)
                if lab:
                    return lab
            except Exception:
                pass
    return ""


@lru_cache(maxsize=10000)
def best_admin_label_cached(lat_round, lon_round, lang="es") -> str:
    return best_admin_label(lat_round, lon_round, lang=lang)


def _shorten_plus(code: str, lat: float, lon: float) -> Optional[str]:
    try:
        if code and pd.notna(lat) and pd.notna(lon):
            sh = olc.shorten(code, float(lat), float(lon))
            if sh and "+" in sh:
                return sh
    except Exception:
        pass
    return None


def pluscode_label(code: Optional[str], lat: Optional[float], lon: Optional[float], lang="es") -> str:
    if (lat is None) or (lon is None) or pd.isna(lat) or pd.isna(lon):
        return (code or "").strip()
    label = best_admin_label_cached(round(float(lat), COORD_PRECISION_CACHE), round(float(lon), COORD_PRECISION_CACHE), lang=lang)
    if not code:
        try:
            code = olc.encode(float(lat), float(lon), codeLength=10)
        except Exception:
            code = None
    short = _shorten_plus(code, lat, lon) if code else None
    token = short or code or ""
    if token and label:
        return f"{token} {label}"
    return token or label or ""

# --------------------------- FIN helpers nuevos ------------------------------
def reverse_best_multi(lat, lon, lang="es") -> str:
    """Cadena de geocodificación:
    1) Nominatim (geopy) vía HOSTS_CHAIN y multi-zoom.
    2) Nominatim HTTP vía URLS_CHAIN.
    3) OpenCage como último fallback.
    (Sin LocationIQ ni BigDataCloud).
    """
    ZOOMS = [18, 17, 16, 15, 14, 12, 10]
    lang_norm = (lang or "es").lower()
    langs = [lang_norm] + [x for x in ("es", "es-mx", "en") if x != lang_norm]

    # 1) geopy/Nominatim por host y zooms
    for L in langs:
        for host in HOSTS_CHAIN:
            try:
                rev = _get_geopy_reverse_for_host(host)
                for Z in ZOOMS:
                    try:
                        res = rev(
                            (lat, lon),
                            language=L,
                            addressdetails=True,
                            zoom=Z,
                            timeout=GEOCODER_TIMEOUT,
                        )
                        if res:
                            if getattr(res, "address", None):
                                return res.address
                            raw = getattr(res, "raw", None)
                            if isinstance(raw, dict) and isinstance(raw.get("address"), dict):
                                composed = _compose_from_addressdict(raw["address"])
                                if composed:
                                    return composed
                    except Exception:
                        pass
            except Exception:
                pass

        # 2) HTTP por base y zooms
        for Z in ZOOMS:
            for base in URLS_CHAIN:
                try:
                    name = _reverse_http_base(base, lat, lon, lang=L, zoom=Z)
                    if name:
                        return name
                except Exception:
                    pass

    # 3) Proveedor secundario ÚNICO: OpenCage
    try:
        name = _reverse_opencage(lat, lon, lang=lang_norm)
        if name:
            return name
    except Exception:
        pass

    return ""


@lru_cache(maxsize=10000)
def reverse_best_cached(lat_round, lon_round, lang="es") -> str:
    return reverse_best_multi(lat_round, lon_round, lang=lang)


def reverse_address(lat, lon, lang="es", precision=COORD_PRECISION_CACHE) -> str:
    lt = round(float(lat), precision)
    ln = round(float(lon), precision)
    return reverse_best_cached(lt, ln, lang=lang)

# ===========================
# PARSEO DE FECHA/HORA
# ===========================
EXCEL_BASE = datetime(1899, 12, 30)  # base Windows


def _excel_serial_to_datetime(val) -> Optional[datetime]:
    try:
        serial = float(val)
    except Exception:
        return None
    if not (0 <= serial < 600000):
        return None
    days = int(serial)
    frac = serial - days
    return EXCEL_BASE + timedelta(days=days) + timedelta(seconds=round(frac * 86400))


_date_pat_ddmmyyyy = re.compile(r"^\s*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\s*$")
_date_pat_yyyymmdd = re.compile(r"^\s*(\d{4})[/-](\d{1,2})[/-](\d{1,2})\s*$")
_time_pat = re.compile(r"^\s*(\d{1,2}):(\d{2})(?::(\d{2}))?\s*([AP]\s*M)?\s*$", re.I)


def _safe_year(y: int) -> int:
    if y < 100:
        return 2000 + y if y <= 69 else 1900 + y
    return y


def _parse_date_str(s: str) -> Optional[date]:
    s = (s or "").strip()
    if not s:
        return None
    m = _date_pat_ddmmyyyy.match(s)
    if m:
        d_, mth, y = int(m.group(1)), int(m.group(2)), _safe_year(int(m.group(3)))
        if 1 <= mth <= 12 and 1 <= d_ <= 31:
            try:
                return date(y, mth, d_)
            except:  # noqa: E722
                return None
    m = _date_pat_yyyymmdd.match(s)
    if m:
        y, mth, d_ = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= mth <= 12 and 1 <= d_ <= 31:
            try:
                return date(y, mth, d_)
            except:  # noqa: E722
                return None
    try:
        as_float = float(s.replace(",", "."))
        dt = _excel_serial_to_datetime(as_float)
        return dt.date() if dt else None
    except:  # noqa: E722
        return None


def _parse_time_str(s: str) -> Optional[dtime]:
    s = (s or "").strip().upper()
    if not s:
        return None
    m = _time_pat.match(s.replace(".", ":"))
    if not m:
        try:
            f = float(s.replace(",", "."))
            if 0 <= f < 1:
                secs = int(round(f * 86400))
                h, r = divmod(secs, 3600)
                mi, se = divmod(r, 60)
                return dtime(hour=min(h, 23), minute=mi, second=se)
            if 0 <= f < 24:
                h = int(f)
                mi = int(round((f - h) * 60))
                return dtime(hour=min(h, 23), minute=min(mi, 59), second=0)
        except:  # noqa: E722
            return None
        return None
    h = int(m.group(1))
    mi = int(m.group(2))
    se = int(m.group(3) or 0)
    ampm = (m.group(4) or "").replace(" ", "")
    if ampm:
        if ampm.upper().startswith("P") and h != 12:
            h += 12
        if ampm.upper().startswith("A") and h == 12:
            h = 0
    if not (0 <= h <= 23 and 0 <= mi <= 59 and 0 <= se <= 59):
        return None
    return dtime(hour=h, minute=mi, second=se)


def build_datetime_from_cols(fecha_val, hora_val) -> Optional[datetime]:
    # A) Fecha ya como datetime/Timestamp
    if isinstance(fecha_val, (pd.Timestamp, datetime)):
        d_ = fecha_val.date()
        if isinstance(hora_val, (pd.Timestamp, datetime)):
            return datetime.combine(d_, hora_val.time())
        if isinstance(hora_val, dtime):
            return datetime.combine(d_, hora_val)
        if isinstance(hora_val, (int, float)) and not pd.isna(hora_val):
            f = float(hora_val)
            if 0 <= f < 1:
                secs = int(round(f * 86400))
                h, r = divmod(secs, 3600)
                mi, se = divmod(r, 60)
                return datetime(d_.year, d_.month, d_.day, h, mi, se)
            if 0 <= f < 24:
                h = int(f)
                mi = int(round((f - h) * 60))
                return datetime(d_.year, d_.month, d_.day, min(h, 23), min(mi, 59), 0)
        if isinstance(hora_val, str) and hora_val.strip():
            tt = _parse_time_str(hora_val)
            if tt:
                return datetime(d_.year, d_.month, d_.day, tt.hour, tt.minute, tt.second)
        return datetime(d_.year, d_.month, d_.day, 0, 0, 0)

    # B) Fecha como date
    if isinstance(fecha_val, date):
        d_ = fecha_val
        if isinstance(hora_val, (pd.Timestamp, datetime)):
            return datetime.combine(d_, hora_val.time())
        if isinstance(hora_val, dtime):
            return datetime.combine(d_, hora_val)
        if isinstance(hora_val, (int, float)) and not pd.isna(hora_val):
            f = float(hora_val)
            if 0 <= f < 1:
                secs = int(round(f * 86400))
                h, r = divmod(secs, 3600)
                mi, se = divmod(r, 60)
                return datetime(d_.year, d_.month, d_.day, h, mi, se)
            if 0 <= f < 24:
                h = int(f)
                mi = int(round((f - h) * 60))
                return datetime(d_.year, d_.month, d_.day, min(h, 23), min(mi, 59), 0)
        if isinstance(hora_val, str) and hora_val.strip():
            tt = _parse_time_str(hora_val)
            if tt:
                return datetime(d_.year, d_.month, d_.day, tt.hour, tt.minute, tt.second)
        return datetime(d_.year, d_.month, d_.day, 0, 0, 0)

    # C) Fecha numérica (serial Excel) o texto
    if isinstance(fecha_val, (int, float)) and not pd.isna(fecha_val):
        base = _excel_serial_to_datetime(fecha_val)
        if base is None:
            return None
        if isinstance(hora_val, (int, float)) and not pd.isna(hora_val):
            f = float(hora_val)
            if 0 <= f < 1:
                return base + timedelta(seconds=round(f * 86400))
            if 0 <= f < 24:
                h = int(f)
                mi = int(round((f - h) * 60))
                return datetime.combine(base.date(), dtime(hour=min(h, 23), minute=min(mi, 59)))
        elif isinstance(hora_val, str):
            tt = _parse_time_str(hora_val)
            return datetime.combine(base.date(), tt) if tt else base
        return base

    if isinstance(fecha_val, str) and fecha_val.strip():
        dd = _parse_date_str(fecha_val)
        if dd:
            if isinstance(hora_val, (int, float)) and not pd.isna(hora_val):
                f = float(hora_val)
                if 0 <= f < 1:
                    secs = int(round(f * 86400))
                    h, r = divmod(secs, 3600)
                    mi, se = divmod(r, 60)
                    return datetime(dd.year, dd.month, dd.day, h, mi, se)
                if 0 <= f < 24:
                    h = int(f)
                    mi = int(round((f - h) * 60))
                    return datetime(dd.year, dd.month, dd.day, min(h, 23), min(mi, 59), 0)
            if isinstance(hora_val, str) and hora_val.strip():
                tt = _parse_time_str(hora_val)
                if tt:
                    return datetime(dd.year, dd.month, dd.day, tt.hour, tt.minute, tt.second)
            return datetime(dd.year, dd.month, dd.day, 0, 0, 0)

    return None


def build_datetime_from_single(val) -> Optional[datetime]:
    if isinstance(val, (pd.Timestamp, datetime)):
        return val.to_pydatetime() if isinstance(val, pd.Timestamp) else val
    if isinstance(val, (int, float)) and not pd.isna(val):
        return _excel_serial_to_datetime(val)
    s = str(val or "").strip()
    if not s:
        return None
    dd = _parse_date_str(s)
    hm = re.search(r"(\d{1,2}:\d{2}(?::\d{2})?\s*(?:[AP]\s*M)?)", s, re.I)
    tt = _parse_time_str(hm.group(1)) if hm else None
    if dd and tt:
        return datetime(dd.year, dd.month, dd.day, tt.hour, tt.minute, tt.second)
    if dd:
        return datetime(dd.year, dd.month, dd.day, 0, 0, 0)
    try:
        f = float(s.replace(",", "."))
        return _excel_serial_to_datetime(f)
    except:  # noqa: E722
        return None

# ===========================
# ESTADÍSTICAS
# ===========================
def build_statistics_blocks(df: pd.DataFrame) -> List[Tuple[str, pd.DataFrame]]:
    blocks: List[Tuple[str, pd.DataFrame]] = []

    tel_col = "Teléfono" if "Teléfono" in df.columns else None
    numA, numB = df.get("Número A"), df.get("Número B")
    tipo = df.get("Tipo")
    dur = (
        pd.to_numeric(df.get("Duración (seg)"), errors="coerce")
        if "Duración (seg)" in df.columns
        else pd.Series(dtype=float)
    )
    dt = df.get("Datetime")

    if tel_col:
        tel_clean_series = df[tel_col].map(_clean_number)
        tel = tel_clean_series.mode().iloc[0] if not tel_clean_series.dropna().empty else ""
    else:
        tel = ""

    A_clean = numA.map(_clean_number) if numA is not None else pd.Series([""] * len(df))
    B_clean = numB.map(_clean_number) if numB is not None else pd.Series([""] * len(df))
    is_voice = (
        tipo.astype(str).str.contains("VOZ", case=False, na=False)
        if tipo is not None
        else pd.Series([False] * len(df))
    )

    if isinstance(dt, pd.Series) and pd.api.types.is_datetime64_any_dtype(dt):
        dt_valid = dt.dropna().sort_values()
    else:
        dt_valid = pd.Series(dtype="datetime64[ns]")

    rango_ini = dt_valid.min() if not dt_valid.empty else pd.NaT
    rango_fin = dt_valid.max() if not dt_valid.empty else pd.NaT
    dias_act = dt_valid.dt.date.nunique() if not dt_valid.empty else 0
    max_gap_h = (dt_valid.diff().max().total_seconds() / 3600) if (len(dt_valid) >= 2) else 0
    counts_by_tipo = tipo.value_counts().to_dict() if tipo is not None else {}
    voz_dur = dur[is_voice] if not dur.empty else pd.Series(dtype=float)

    resumen_rows = [
        ["Teléfono investigado", tel],
        ["Rango temporal - inicio", str(rango_ini) if pd.notna(rango_ini) else ""],
        ["Rango temporal - fin", str(rango_fin) if pd.notna(rango_fin) else ""],
        ["Días activos", dias_act],
        ["Máximo periodo de inactividad (h)", round(max_gap_h, 2)],
        ["Eventos totales", len(df)],
    ]
    for k, v in counts_by_tipo.items():
        resumen_rows.append([f"Total {k}", int(v)])
    if not voz_dur.empty:
        resumen_rows += [
            ["Duración VOZ total (s)", int(voz_dur.sum())],
            ["Duración VOZ media (s)", round(float(voz_dur.mean()), 2)],
            ["Duración VOZ mediana (s)", round(float(voz_dur.median()), 2)],
            ["Duración VOZ p95 (s)", round(float(voz_dur.quantile(0.95)), 2)],
        ]
    blocks.append(("1) Resumen global", pd.DataFrame(resumen_rows, columns=["Métrica", "Valor"])))

    out_mask = (A_clean == tel) & is_voice
    in_mask = (B_clean == tel) & is_voice
    dir_rows = []
    for nombre, m in [("SALIENTE", out_mask), ("ENTRANTE", in_mask)]:
        n = int(m.sum())
        dsum = int(dur[m].sum()) if not dur.empty else 0
        dmean = round(float(dur[m].mean()), 2) if n > 0 else 0.0
        dir_rows.append([nombre, n, dsum, dmean])
    blocks.append(
        (
            "2) Dirección del tráfico (VOZ)",
            pd.DataFrame(dir_rows, columns=["Dirección", "# Llamadas", "Duración total (s)", "Duración media (s)"]),
        )
    )

    if numB is not None:
        sal = df[out_mask].copy()
        if not sal.empty:
            sal = sal.assign(Contraparte=B_clean[out_mask].values)
            if "Duración (seg)" in sal.columns:
                sal["Duración (seg)"] = pd.to_numeric(sal["Duración (seg)"], errors="coerce")
            top_sal = (
                sal.groupby("Contraparte", dropna=True)
                .agg(
                    Llamadas=("Datetime", "count"),
                    Duración_total_s=("Duración (seg)", "sum"),
                    Primera=("Datetime", "min"),
                    Última=("Datetime", "max"),
                )
                .reset_index()
                .rename(columns={"Contraparte": "Número"})
            )
            top_sal = top_sal.sort_values(["Llamadas", "Duración_total_s"], ascending=[False, False]).head(10)
            blocks.append(("3) Top 10 contactos SALIENTES", top_sal))
        else:
            blocks.append(
                (
                    "3) Top 10 contactos SALIENTES",
                    pd.DataFrame(columns=["Número", "Llamadas", "Duración_total_s", "Primera", "Última"]),
                )
            )
    else:
        blocks.append(
            (
                "3) Top 10 contactos SALIENTES",
                pd.DataFrame(columns=["Número", "Llamadas", "Duración_total_s", "Primera", "Última"]),
            )
        )

    if numA is not None:
        ent = df[in_mask].copy()
        if not ent.empty:
            ent = ent.assign(Contraparte=A_clean[in_mask].values)
            if "Duración (seg)" in ent.columns:
                ent["Duración (seg)"] = pd.to_numeric(ent["Duración (seg)"], errors="coerce")
            top_ent = (
                ent.groupby("Contraparte", dropna=True)
                .agg(
                    Llamadas=("Datetime", "count"),
                    Duración_total_s=("Duración (seg)", "sum"),
                    Primera=("Datetime", "min"),
                    Última=("Datetime", "max"),
                )
                .reset_index()
                .rename(columns={"Contraparte": "Número"})
            )
            top_ent = top_ent.sort_values(["Llamadas", "Duración_total_s"], ascending=[False, False]).head(10)
            blocks.append(("4) Top 10 contactos ENTRANTES", top_ent))
        else:
            blocks.append(
                (
                    "4) Top 10 contactos ENTRANTES",
                    pd.DataFrame(columns=["Número", "Llamadas", "Duración_total_s", "Primera", "Última"]),
                )
            )
    else:
        blocks.append(
            (
                "4) Top 10 contactos ENTRANTES",
                pd.DataFrame(columns=["Número", "Llamadas", "Duración_total_s", "Primera", "Última"]),
            )
        )

    # Antenas TOP por tipo → preferir 'Direccion_final'
    if "Direccion_final" in df.columns and df["Direccion_final"].notna().any():
        antena_series = df["Direccion_final"].fillna("")
    elif "PLUS_CODE_NOMBRE" in df.columns and df["PLUS_CODE_NOMBRE"].notna().any():
        antena_series = df["PLUS_CODE_NOMBRE"].fillna("")
    elif {"Latitud", "Longitud"}.issubset(df.columns):
        antena_series = df["Latitud"].round(6).astype(str) + "," + df["Longitud"].round(6).astype(str)
    else:
        antena_series = pd.Series([""] * len(df))

    tipos_orden = ["DATOS", "VOZ ENTRANTE", "VOZ SALIENTE", "MENSAJES 2 VÍAS", "TRANSFER"]
    letras = ["A", "B", "C", "D", "E"]
    idx_letra = 0
    if "Tipo" in df.columns:
        tipo_norm = df["Tipo"].astype(str).str.upper()
        for tnombre in tipos_orden:
            mask_t = tipo_norm.eq(tnombre)
            if not mask_t.any():
                continue
            tmp = pd.DataFrame(
                {
                    "Antena": antena_series[mask_t],
                    "Datetime": dt[mask_t] if isinstance(dt, pd.Series) else pd.Series([pd.NaT] * mask_t.sum()),
                }
            )
            tmp = tmp[tmp["Antena"].astype(str).str.len() > 0]
            if tmp.empty:
                continue
            top_ant_t = (
                tmp.groupby("Antena")
                .agg(
                    Eventos=("Antena", "count"),
                    Primera=("Datetime", "min"),
                    Última=("Datetime", "max"),
                )
                .reset_index()
                .sort_values("Eventos", ascending=False)
                .head(5)
            )
            titulo = f"5{letras[idx_letra]}) Antenas TOP — {tnombre}"
            blocks.append((titulo, top_ant_t))
            idx_letra += 1

    # IMEI
    if "IMEI" in df.columns:
        imei_df = df.copy()
        imei_df["IMEI"] = imei_df["IMEI"].astype(str)
        imei_df = imei_df[imei_df["IMEI"].str.len() > 0]
        if not imei_df.empty:
            imei_tb = (
                imei_df.groupby("IMEI")
                .agg(
                    Eventos=("IMEI", "count"),
                    Primera=("Datetime", "min"),
                    Última=("Datetime", "max"),
                )
                .reset_index()
                .sort_values("Eventos", ascending=False)
            )
        else:
            imei_tb = pd.DataFrame(columns=["IMEI", "Eventos", "Primera", "Última"])
    else:
        imei_tb = pd.DataFrame(columns=["IMEI", "Eventos", "Primera", "Última"])
    blocks.append(("6) IMEI (uso por periodo)", imei_tb))

    # === Estadística de SMS ===
    if tipo is not None:
        tipo_up = tipo.astype(str).str.upper()
        sms_mask = tipo_up.str.contains("SMS", na=False) | tipo_up.str.contains("MENSAJE", na=False)
    else:
        sms_mask = pd.Series([False] * len(df))

    if sms_mask.any():
        out_sms = (A_clean == tel) & sms_mask
        in_sms = (B_clean == tel) & sms_mask

        env = df[out_sms].copy().assign(Contraparte=B_clean[out_sms].values)
        rec = df[in_sms].copy().assign(Contraparte=A_clean[in_sms].values)
        sms_all = pd.concat([env, rec], ignore_index=True)

        resumen_sms = pd.DataFrame(
            [
                {
                    "Total SMS": int(len(sms_all)),
                    "Contactos con SMS (Número)": int(sms_all["Contraparte"].nunique())
                    if not sms_all.empty
                    else 0,
                    "Primer SMS": sms_all["Datetime"].min() if not sms_all.empty else pd.NaT,
                    "Último SMS": sms_all["Datetime"].max() if not sms_all.empty else pd.NaT,
                }
            ]
        )
        blocks.append(("7) Mensajes SMS — Resumen", resumen_sms))

        if not sms_all.empty:
            top_tot = (
                sms_all.groupby("Contraparte", dropna=True)
                .agg(
                    SMS_totales=("Datetime", "count"),
                    Primero=("Datetime", "min"),
                    Último=("Datetime", "max"),
                )
                .reset_index()
                .rename(columns={"Contraparte": "Número"})
                .sort_values(["SMS_totales"], ascending=[False])
                .head(10)
            )
        else:
            top_tot = pd.DataFrame(columns=["Número", "SMS_totales", "Primero", "Último"])
        blocks.append(("8) Top 10 contactos por SMS (totales)", top_tot))

        if not env.empty:
            top_env = (
                env.groupby("Contraparte", dropna=True)
                .agg(
                    SMS_enviados=("Datetime", "count"),
                    Primero=("Datetime", "min"),
                    Último=("Datetime", "max"),
                )
                .reset_index()
                .rename(columns={"Contraparte": "Número"})
                .sort_values(["SMS_enviados"], ascending=[False])
                .head(10)
            )
        else:
            top_env = pd.DataFrame(columns=["Número", "SMS_enviados", "Primero", "Último"])
        blocks.append(("9) Top 10 SMS enviados", top_env))

        if not rec.empty:
            top_rec = (
                rec.groupby("Contraparte", dropna=True)
                .agg(
                    SMS_recibidos=("Datetime", "count"),
                    Primero=("Datetime", "min"),
                    Último=("Datetime", "max"),
                )
                .reset_index()
                .rename(columns={"Contraparte": "Número"})
                .sort_values(["SMS_recibidos"], ascending=[False])
                .head(10)
            )
        else:
            top_rec = pd.DataFrame(columns=["Número", "SMS_recibidos", "Primero", "Último"])
        blocks.append(("10) Top 10 SMS recibidos", top_rec))

        if not sms_all.empty:
            tmp = sms_all.copy()
            tmp["Fecha_dia"] = pd.to_datetime(tmp["Datetime"]).dt.date
            por_dia = tmp.groupby("Fecha_dia").size().reset_index(name="SMS_dia")
        else:
            por_dia = pd.DataFrame(columns=["Fecha_dia", "SMS_dia"])
        blocks.append(("11) SMS por día", por_dia))

    return blocks

# ===========================
# DEDUPE SOLO “DATOS” (minuto)
# ===========================
def dedupe_datos_by_minute(df: pd.DataFrame):
    if df is None or df.empty or "Tipo" not in df.columns:
        return df, pd.DataFrame(), 0
    if not {"Número A", "Número B"}.issubset(df.columns):
        return df, pd.DataFrame(), 0

    tipo_up = df["Tipo"].astype(str).str.upper()
    mask = tipo_up.eq("DATOS")
    otros = df.loc[~mask].copy()
    datos = df.loc[mask].copy()

    use_dt = False
    if "Datetime" in datos.columns:
        datos["Datetime"] = pd.to_datetime(datos["Datetime"], errors="coerce")
        use_dt = datos["Datetime"].notna().any()

    if use_dt:
        datos["__t__"] = datos["Datetime"].dt.floor("min")
    else:
        if {"Fecha", "Hora"}.issubset(datos.columns):
            tmp_keys = []
            for f_val, h_val in zip(datos["Fecha"], datos["Hora"]):
                dt_ = build_datetime_from_cols(f_val, h_val)
                tmp_keys.append(pd.to_datetime(dt_) if dt_ else pd.NaT)
            key_series = pd.to_datetime(pd.Series(tmp_keys), errors="coerce")
            if key_series.notna().any():
                datos["__t__"] = key_series.dt.floor("min")
                use_dt = True
            else:
                return df, pd.DataFrame(), 0
        else:
            return df, pd.DataFrame(), 0

    if "Duración (seg)" in datos.columns:
        datos["Duración (seg)"] = pd.to_numeric(datos["Duración (seg)"], errors="coerce")
        datos = datos.sort_values("Duración (seg)", ascending=False)

    dup_extras = datos.duplicated(subset=["Número A", "Número B", "__t__"], keep="first")
    duplicados_df = datos[dup_extras].copy()
    datos = datos[~dup_extras].drop(columns=["__t__"], errors="ignore")
    out = pd.concat([otros, datos], ignore_index=True)
    return out, duplicados_df, int(dup_extras.sum())

# ===========================
# LÓGICA PRINCIPAL
# ===========================
def limpiar_excel(file, remove_duplicates: bool = False, offline: bool = False):
    progress = st.progress(0, text="Iniciando…")
    progress_section(progress, 4, "📥 Cargando archivo…")

    if not offline:
        init_plus_repo()

    df = leer_archivo(file)
    original_len = len(df)

    # 1) Encabezado por 'Teléfono'
    progress_section(progress, 8, "🔎 Buscando encabezados…")
    start_row = None
    for i, row in enumerate(df.values):
        if any(norm(cell) == "telefono" for cell in row):
            start_row = i
            break
    if start_row is not None:
        df.columns = df.iloc[start_row].astype(str).str.strip()
        df = df.iloc[start_row + 1 :].reset_index(drop=True)
    else:
        st.warning("⚠️ No se encontró la columna 'Teléfono'; se procesará desde la primera fila.")

    # 2) Renombrado flexible
    progress_section(progress, 12, "🧩 Normalizando columnas…")
    mapa = {
        r"^tel$|^telefono$": "Teléfono",
        r"^tipo$": "Tipo",
        r"^num.*a$": "Número A",
        r"^num.*b$": "Número B",
        r"^fecha$": "Fecha",
        r"^hora$": "Hora",
        r"^dur|^duracion": "Duración (seg)",
        r"^imei$": "IMEI",
        r"^lat|^latitud": "Latitud",
        r"^lon$|^long$|^longitud": "Longitud",
        r"^azim|^azimuth": "Azimuth",
    }
    ren = {}
    for c in df.columns:
        nc = norm(c)
        nuevo = c
        for patron, nombre in mapa.items():
            if re.search(patron, nc, re.I):
                nuevo = nombre
                break
        ren[c] = nuevo
    df = df.rename(columns=ren)

    # 3) Resguardo originales
    progress_section(progress, 16, "📝 Resguardando originales…")
    for col in ["Latitud", "Longitud", "Azimuth", "Fecha", "Hora", "FechaHora", "Fecha y Hora"]:
        if col in df.columns:
            df[f"{col}_raw"] = df[col].astype(str)

    # 4) Asegurar Lat/Lon y convertir
    for col in ["Latitud", "Longitud"]:
        if col not in df.columns:
            df[col] = None
    df = ensure_lat_lon_columns(df)

    progress_section(progress, 20, "🧭 Convirtiendo coordenadas a decimal…")
    for col in ["Latitud", "Longitud"]:
        if col in df.columns:
            df[col] = df[col].apply(dms_to_decimal)
    if "Longitud" in df.columns:
        df["Longitud"] = df["Longitud"].apply(lambda x: -abs(x) if pd.notna(x) else x)

    df = force_latlon_from_any(df)

    # 5) Columnas de salida
    if "PLUS_CODE" not in df.columns:
        df["PLUS_CODE"] = None
    if "PLUS_CODE_NOMBRE" not in df.columns:
        df["PLUS_CODE_NOMBRE"] = None
    if "PLUS_CODE_SHORT" not in df.columns:
        df["PLUS_CODE_SHORT"] = None

    # 6) PLUS_CODE offline y PlusRepo
    progress_section(progress, 28, "➕ Generando Plus Codes…")
    mask_coords = df["Latitud"].notna() & df["Longitud"].notna()
    if mask_coords.any():
        def generar_pluscode(lat, lon):
            try:
                return olc.encode(float(lat), float(lon), codeLength=10)
            except Exception:
                return None

        df.loc[mask_coords, "PLUS_CODE"] = df.loc[mask_coords, "PLUS_CODE"].where(
            df["PLUS_CODE"].notna(),
            df.loc[mask_coords].apply(lambda x: generar_pluscode(x["Latitud"], x["Longitud"]), axis=1),
        )

        def _mk_short(row):
            return _shorten_plus(row.get("PLUS_CODE"), row.get("Latitud"), row.get("Longitud"))

        df.loc[mask_coords, "PLUS_CODE_SHORT"] = df.loc[mask_coords].apply(_mk_short, axis=1)

        name_map: Dict[str, str] = {}
        for code in ([] if offline else pd.Series(df.loc[mask_coords, "PLUS_CODE"].dropna().unique())):
            try:
                la, lo, clen = _decode_plus(code)
                pr_save_plus(code, la, lo, clen)
                found = pr_get_nombre(code, la, lo)
                if found:
                    name_map[code] = found[0]
            except Exception:
                pass
        if name_map:
            empty_mask = df["PLUS_CODE_NOMBRE"].isna() | (df["PLUS_CODE_NOMBRE"].astype(str).str.strip() == "")
            df.loc[empty_mask, "PLUS_CODE_NOMBRE"] = df.loc[empty_mask, "PLUS_CODE"].map(name_map)

    # 7) Geocoding multi-zoom + near repo + cache
    progress_section(progress, 54, "🔒 Modo offline: sin consultas de dirección" if offline else "🌍 Geocodificando…")
    if not offline and GEOCODE_ENABLED and mask_coords.any():
        with closing(sqlite3.connect(CACHE_DB_PATH)) as con:
            con.execute(
                """CREATE TABLE IF NOT EXISTS geocache (
                lat REAL NOT NULL, lon REAL NOT NULL, addr TEXT, PRIMARY KEY(lat, lon)
            )"""
            )
            con.execute("PRAGMA journal_mode=WAL;")
            con.commit()

        coords_unique = df.loc[mask_coords, ["Latitud", "Longitud"]].drop_duplicates().reset_index(drop=True)
        keys = [
            (round(float(r["Latitud"]), COORD_PRECISION_CACHE), round(float(r["Longitud"]), COORD_PRECISION_CACHE))
            for _, r in coords_unique.iterrows()
        ]

        if keys:
            q_marks = ",".join(["(?,?)"] * len(keys))
            params = [v for pair in keys for v in pair]
            with closing(sqlite3.connect(CACHE_DB_PATH)) as con:
                rows = con.execute(
                    f"SELECT lat, lon, COALESCE(addr,'') FROM geocache WHERE (lat,lon) IN ({q_marks})",
                    params,
                ).fetchall()
            ubic_map: Dict[Tuple[float, float], str] = {(r[0], r[1]): r[2] for r in rows}
        else:
            ubic_map = {}

        missing = [k for k in keys if k not in ubic_map]
        still_missing = []
        for (lt, ln) in missing:
            v = pr_get_near_name(lt, ln)
            if v:
                ubic_map[(lt, ln)] = v
            else:
                still_missing.append((lt, ln))

        to_resolve: List[Tuple[float, float]] = still_missing[: int(MAX_UNIQUE_GEOCODES)]
        new_inserts: Dict[Tuple[float, float], str] = {}
        resueltos = sum(1 for k in keys if k in ubic_map and ubic_map[k])
        fallidos = 0

        for idx, (lt, ln) in enumerate(to_resolve, 1):
            addr = reverse_address(lt, ln, lang="es", precision=COORD_PRECISION_CACHE)
            if addr:
                resueltos += 1
            else:
                fallidos += 1
                addr = "SIN_DIRECCIÓN"
            ubic_map[(lt, ln)] = addr
            new_inserts[(lt, ln)] = addr

            pct = 54 + int(idx / max(len(to_resolve), 1) * (84 - 54))
            progress_section(progress, pct, f"🌍 {idx}/{len(to_resolve)} — resueltas: {resueltos}")
            time.sleep(0.2)

        if new_inserts:
            with closing(sqlite3.connect(CACHE_DB_PATH)) as con:
                con.executemany(
                    "INSERT OR REPLACE INTO geocache(lat,lon,addr) VALUES(?,?,?)",
                    [(k[0], k[1], v or "") for k, v in new_inserts.items()],
                )
                con.commit()

        def pick_name(lat, lon, curr):
            if pd.notna(curr) and str(curr).strip():
                return curr
            k = (round(float(lat), COORD_PRECISION_CACHE), round(float(lon), COORD_PRECISION_CACHE))
            v = ubic_map.get(k, "")
            return v if v else "SIN_DIRECCIÓN"

        df.loc[mask_coords, "PLUS_CODE_NOMBRE"] = df.loc[mask_coords].apply(
            lambda x: pick_name(x["Latitud"], x["Longitud"], x["PLUS_CODE_NOMBRE"]),
            axis=1,
        )

        if "PLUS_CODE" in df.columns:
            new_keys = set(new_inserts.keys())
            for code, g in df.loc[mask_coords].groupby("PLUS_CODE", dropna=True):
                try:
                    lat0 = float(g["Latitud"].iloc[0])
                    lon0 = float(g["Longitud"].iloc[0])
                    k0 = (round(lat0, COORD_PRECISION_CACHE), round(lon0, COORD_PRECISION_CACHE))
                    nombre = ubic_map.get(k0, "")
                    if nombre:
                        fuente = "geocoder" if k0 in new_keys else "cache"
                        conf = 0.9 if fuente == "geocoder" else 0.95
                        pr_save_nombre(code, nombre, fuente=fuente, conf=conf, lat=lat0, lon=lon0)
                except Exception:
                    pass

        st.info(
            f"Geocoding → únicas: {len(coords_unique):,} | consultadas: {len(to_resolve):,} | "
            f"resueltas: {resueltos:,} | fallidas: {fallidos:,} | hits caché/near: {len(ubic_map) - len(to_resolve):,}"
        )

    # 7.bis) Dirección final (con PLUS corto + admin label)
    def _is_sin_dir(texto: str) -> bool:
        t = (texto or "").upper()
        return (t == "") or ("SIN_DIRECCIÓN" in t) or ("SIN_DIRECCION" in t) or t.startswith("CERCA DE (")

    dir_map: Dict[Tuple[float, float], str] = {}
    if not offline and GEOCODE_ENABLED and mask_coords.any():
        need_dir_mask = mask_coords & df["PLUS_CODE_NOMBRE"].apply(lambda v: _is_sin_dir(str(v or "")))
        if need_dir_mask.any():
            unique_dir = df.loc[need_dir_mask, ["Latitud", "Longitud", "PLUS_CODE"]].drop_duplicates().reset_index(drop=True)
            for _, r in unique_dir.iterrows():
                lat = r.get("Latitud", None)
                lon = r.get("Longitud", None)
                code = r.get("PLUS_CODE", None)
                if pd.notna(lat) and pd.notna(lon):
                    k = (round(float(lat), COORD_PRECISION_CACHE), round(float(lon), COORD_PRECISION_CACHE))
                    composed = pluscode_label(code, lat, lon, lang="es")
                    if composed:
                        dir_map[k] = composed
                        try:
                            pr_save_nombre(code or "", composed, fuente="composed", conf=0.85, lat=lat, lon=lon)
                        except Exception:
                            pass
                    else:
                        try:
                            dir_map[k] = f"Cerca de ({float(lat):.6f}, {float(lon):.6f})"
                        except Exception:
                            dir_map[k] = "SIN_DIRECCIÓN"

    def direccion_fallback(row):
        nombre = str(row.get("PLUS_CODE_NOMBRE", "") or "").strip()
        lat = row.get("Latitud", None)
        lon = row.get("Longitud", None)

        if not _is_sin_dir(nombre):
            return nombre

        if pd.notna(lat) and pd.notna(lon):
            k = (round(float(lat), COORD_PRECISION_CACHE), round(float(lon), COORD_PRECISION_CACHE))
            if k in dir_map and dir_map[k]:
                return dir_map[k]
            try:
                return f"Cerca de ({float(lat):.6f}, {float(lon):.6f})"
            except Exception:
                return "SIN_DIRECCIÓN"
        return "SIN_DIRECCIÓN"

    df["Direccion_final"] = df.apply(offline_address if offline else direccion_fallback, axis=1)

    # 8) Azimuth (parser robusto)
    progress_section(progress, 86, "🧮 Normalizando azimuth…")
    if "Azimuth" in df.columns:
        df["Azimuth_deg"] = df["Azimuth"].apply(parse_azimuth)

    # 9) Datetime — parser estricto
    progress_section(progress, 90, "⏱️ Calculando Datetime…")
    df["Datetime"] = pd.NaT

    if {"Fecha", "Hora"}.issubset(df.columns):
        dt_list = []
        for f_val, h_val in zip(df["Fecha"].to_list(), df["Hora"].to_list()):
            dt_ = build_datetime_from_cols(f_val, h_val)
            dt_list.append(dt_)
        df["Datetime"] = pd.to_datetime(pd.Series(dt_list), errors="coerce")
    elif "FechaHora" in df.columns:
        dt_list = []
        for v in df["FechaHora"].to_list():
            dt_ = build_datetime_from_single(v)
            dt_list.append(dt_)
        df["Datetime"] = pd.to_datetime(pd.Series(dt_list), errors="coerce")
    elif "Fecha y Hora" in df.columns:
        dt_list = []
        for v in df["Fecha y Hora"].to_list():
            dt_ = build_datetime_from_single(v)
            dt_list.append(dt_)
        df["Datetime"] = pd.to_datetime(pd.Series(dt_list), errors="coerce")
    else:
        st.warning("⚠️ No se encontraron columnas de fecha/hora reconocibles (Fecha+Hora o FechaHora).")

    # 10) Tipo
    progress_section(progress, 92, "📚 Normalizando tipo…")
    if "Tipo" in df.columns:
        rep = {
            "voz entrante": "VOZ ENTRANTE",
            "entrante": "VOZ ENTRANTE",
            "voz saliente": "VOZ SALIENTE",
            "saliente": "VOZ SALIENTE",
            "datos": "DATOS",
            "transfer": "TRANSFER",
            "mensajes 2 vías": "MENSAJES 2 VÍAS",
            "2 vías": "MENSAJES 2 VÍAS",
        }
        df["Tipo"] = df["Tipo"].astype(str).str.strip().str.lower().replace(rep).str.upper()

    # 11) Duplicados informativo + eliminación real SOLO DATOS
    progress_section(
        progress,
        94,
        ("🧽 Eliminando duplicados (DATOS)…" if remove_duplicates else "🔎 Buscando duplicados…"),
    )
    duplicados_df = pd.DataFrame()
    eliminados = 0

    subset_general = [c for c in ["Número A", "Número B", "Datetime"] if c in df.columns]
    if subset_general:
        dup_mask_all = df.duplicated(subset=subset_general, keep=False)
        df["Es_Duplicado"] = dup_mask_all
        key_col = subset_general[0]
        try:
            df["Cuenta_GrupoDup"] = df.groupby(subset_general, dropna=False)[key_col].transform("size")
        except TypeError:
            df["Cuenta_GrupoDup"] = df.groupby(subset_general)[key_col].transform("size")
    else:
        df["Es_Duplicado"] = False
        df["Cuenta_GrupoDup"] = 1

    if remove_duplicates and {"Número A", "Número B"}.issubset(df.columns):
        df, duplicados_df, eliminados = dedupe_datos_by_minute(df)
        st.caption(f"Duplicados 'DATOS' eliminados: {eliminados}")

    # 12) Orden final
    if "Datetime" in df.columns:
        df = (
            df.sort_values(
                ["Datetime", "Número A", "Número B"],
                ascending=[True, True, True],
                na_position="last",
            )
            .reset_index(drop=True)
        )
    if not duplicados_df.empty and "Datetime" in duplicados_df.columns:
        duplicados_df = (
            duplicados_df.sort_values(
                ["Datetime", "Número A", "Número B"],
                ascending=[True, True, True],
                na_position="last",
            )
            .reset_index(drop=True)
        )

    # 13) LOG / métricas
    coords_validas = int(
        (df.get("Latitud", pd.Series(dtype=float)).notna() & df.get("Longitud", pd.Series(dtype=float)).notna()).sum()
    )
    plus_generados = int(df.get("PLUS_CODE", pd.Series(dtype=object)).notna().sum())
    nombres_generados = int(df.get("PLUS_CODE_NOMBRE", pd.Series(dtype=object)).notna().sum())
    dt_fails = int(df["Datetime"].isna().sum())
    log_df = pd.DataFrame(
        {
            "Filas originales": [original_len],
            "Duplicados eliminados (solo DATOS)": [eliminados],
            "Coordenadas válidas detectadas": [coords_validas],
            "PLUS_CODE generados": [plus_generados],
            "PLUS_CODE_NOMBRE generados": [nombres_generados],
            "Datetime sin parsear": [dt_fails],
            "PlusRepo DB": [PLUS_REPO_DB],
        }
    )

    if offline:
        log_df["Modo geocodificación"] = "OFFLINE: sin consultas externas ni búsqueda en caché de direcciones"

    # 14) ESTADISTICAS
    progress_section(progress, 96, "📊 Calculando estadísticas…")
    stat_blocks = build_statistics_blocks(df)

    # 15) Exportar Excel
    progress_section(progress, 98, "📦 Exportando a Excel…")
    title_rows: List[Tuple[int, str]] = []
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Datos_Limpios", index=False)
        log_df.to_excel(writer, sheet_name="LOG_Limpieza", index=False)
        if not duplicados_df.empty:
            duplicados_df.to_excel(writer, sheet_name="Duplicados", index=False)

        start = 0
        for title, bdf in stat_blocks:
            bdf = bdf.copy()
            bdf.to_excel(
                writer,
                sheet_name="ESTADISTICAS",
                index=False,
                startrow=start + 1,
                startcol=0,
            )
            title_rows.append((start + 1, title))
            start = start + 1 + 1 + len(bdf) + 1

    output.seek(0)
    wb = load_workbook(output)

    ws = wb["Datos_Limpios"]
    # ✅ Número A, Número B e IMEI como números simples sin decimales
    cols_num_fmt: List[int] = []
    for col_name in ["Número A", "Número B", "IMEI"]:
        if col_name in df.columns:
            cols_num_fmt.append(df.columns.get_loc(col_name) + 1)
    for col_idx in cols_num_fmt:
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                cell.number_format = "0"

    for col_idx, column in enumerate(ws.columns, start=1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in column:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)
    tab = Table(
        displayName="TablaLimpia",
        ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}",
    )
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)

    if "LOG_Limpieza" in wb.sheetnames:
        ws_log = wb["LOG_Limpieza"]
        tab_log = Table(
            displayName="TablaLog",
            ref=f"A1:{get_column_letter(ws_log.max_column)}{ws_log.max_row}",
        )
        tab_log.tableStyleInfo = TableStyleInfo(name="TableStyleLight11", showRowStripes=True)
        ws_log.add_table(tab_log)

    if "Duplicados" in wb.sheetnames:
        ws_dup = wb["Duplicados"]
        tab_dup = Table(
            displayName="TablaDuplicados",
            ref=f"A1:{get_column_letter(ws_dup.max_column)}{ws_dup.max_row}",
        )
        tab_dup.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws_dup.add_table(tab_dup)

    if "ESTADISTICAS" in wb.sheetnames:
        ws_est = wb["ESTADISTICAS"]
        for row_idx, title in title_rows:
            ws_est.cell(row=row_idx, column=1).value = title
            ws_est.cell(row=row_idx, column=1).font = Font(bold=True)
        for col in ws_est.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws_est.column_dimensions[col_letter].width = min(max_len + 2, 60)

    # Nombre sugerido por Teléfono
    if "Teléfono" in df.columns:
        tel_clean_series = df["Teléfono"].map(_clean_number)
        telefono_archivo = (
            tel_clean_series.mode().iloc[0]
            if not tel_clean_series.dropna().empty
            else "SIN_TELEFONO"
        )
    else:
        telefono_archivo = "SIN_TELEFONO"
    suggested_name = _sanitize_filename(
        f"{telefono_archivo} LIMPIEZA ESTADISTICO SUITE",
        ext=".xlsx",
    )

    final_output = io.BytesIO()
    wb.save(final_output)
    wb.close()
    final_output.seek(0)
    progress_section(progress, 100, "✅ Listo")
    return final_output.getvalue(), suggested_name

# ===========================
# UI
# ===========================
page_header(
    "Limpieza Automática",
    "Prepara CDR crudos para análisis: coordenadas, azimuth, PLUS_CODE, dirección derivada, deduplicación opcional y estadísticas exportables.",
    eyebrow="PREPARACIÓN DE EVIDENCIA",
    badges=["CSV · XLS · XLSX", "Excel limpio", "Trazabilidad básica"],
)
info_panel("Repositorio de ubicaciones", f"PlusRepo activo: {PLUS_REPO_DB}")

try:
    _olc_demo = olc.encode(19.4326, -99.1332, codeLength=10)
    st.caption(f"Sanity OLC (CDMX): {_olc_demo}")
except Exception as _e:
    st.error(f"[Sanity OLC] Falló: {_e}")

section_title("Carga y opciones", "Selecciona el archivo crudo y activa solo las opciones necesarias para el caso.", "01")
remove_dups = st.checkbox(
    "eliminar celda de Datos duplicada, (el consumo de datos se repite hasta 3,4 veces esto unifica en 1 solo registro)",
    value=False,
    help=(
        "Elimina duplicados SOLO de TIPO 'DATOS' agrupando por {Número A, Número B, minuto}. "
        "Conserva la fila con mayor 'Duración (seg)'."
    ),
)

offline_mode = st.checkbox(
    "🔒 Modo offline: no consultar direcciones en internet",
    value=False,
    help="Omite geocodificación, búsquedas administrativas y repositorios de nombres. Conserva direcciones ya incluidas, coordenadas y calcula Plus Codes localmente.",
)
if offline_mode:
    st.caption("Sin consultas externas. Las direcciones faltantes se identifican como SIN_DIRECCIÓN (MODO OFFLINE).")

uploaded_file = st.file_uploader(
    "📂 Sube tu archivo crudo (CSV, XLS o XLSX)",
    type=["csv", "xls", "xlsx"],
)
if uploaded_file:
    if st.button("🚀 Limpiar y generar nuevo Excel"):
        try:
            cleaned_data, suggested_name = limpiar_excel(uploaded_file, remove_duplicates=remove_dups, offline=offline_mode)
            st.success(
                "✅ Archivo procesado. Revisa 'Datos_Limpios', 'LOG_Limpieza', 'Duplicados' y 'ESTADISTICAS'."
            )
            st.download_button(
                "⬇️ Descargar archivo limpio",
                cleaned_data,
                suggested_name,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except Exception as e:
            st.error(f"⚠️ Error durante la limpieza: {e}")
else:
    st.info("📁 Esperando archivo…")
