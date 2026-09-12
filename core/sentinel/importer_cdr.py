"""Read Datos_Limpios only for events. Original bytes and every source row retained."""
from io import BytesIO
from pathlib import Path
import hashlib,json,re
from collections import Counter
import openpyxl
from .models import ImportOptions,uid,now,dump,stable
from .validators import text,norm,phone,device,temporal,duration,classify,safe_cell,endpoint
from .location_engine import locate
from .entity_engine import ensure_entity

ALIASES={
 'target':('TELEFONO','LINEA OBJETIVO'), 'kind':('TIPO','TIPO EVENTO'),
 'a':('NUMERO A','NUMERO_A','A'), 'b':('NUMERO B','NUMERO_B','B'),
 'datetime':('DATETIME','FECHAHORA','FECHA_HORA'),'date':('FECHA',),'time':('HORA',),
 'duration':('DURACION (SEG)','DURACION','DURACION_SEG'),'imei':('IMEI',),'imsi':('IMSI',),
 'lat':('LATITUD','LAT'),'lon':('LONGITUD','LON'),'azimuth':('AZIMUTH_DEG','AZIMUTH'),
 'plus':('PLUS_CODE',),'address':('DIRECCION_FINAL','PLUS_CODE_NOMBRE','DIRECCION'),
 'cell':('CELL ID','CELL_ID','CELLID'),'lac':('LAC',),'tac':('TAC',)}

def inspect_workbook(name,data):
    if len(data)>100*1024*1024:raise ValueError('Excel excede límite V1 de 100 MiB por archivo')
    w=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=True)
    try:
        if 'Datos_Limpios' not in w.sheetnames:raise ValueError(f'{name}: falta Datos_Limpios')
        s=w['Datos_Limpios']
        if s.max_row>500001:raise ValueError('V1 admite hasta 500,000 filas por libro; dividir lote')
        it=s.iter_rows(values_only=True);headers=[text(v) for v in next(it)]
        normalized=[norm(h) for h in headers]
        if len(normalized)!=len(set(normalized)):raise ValueError('Encabezados duplicados o vacíos ambiguos')
        mapping={k:next((normalized.index(a) for a in aliases if a in normalized),None) for k,aliases in ALIASES.items()}
        if any(mapping[k] is None for k in ('kind','a','b')):raise ValueError('Se requieren Tipo, Número A y Número B')
        if mapping['datetime'] is None and (mapping['date'] is None or mapping['time'] is None):raise ValueError('Falta Datetime o Fecha/Hora')
        rows=list(it)
        candidates=Counter(phone(r[mapping['target']]) for r in rows if mapping['target'] is not None and classify(r[mapping['kind']])!='NO_CLASIFICADO')
        candidates.pop('',None)
        filenames=re.findall(r'(?<!\d)\d{10}(?!\d)',name)
        if len(candidates)==1:target=next(iter(candidates));method='COLUMNA_TELEFONO'
        elif candidates:raise ValueError(f'{name}: más de una línea en Teléfono; revisar fuente')
        elif len(filenames)==1:target=filenames[0];method='NOMBRE_ARCHIVO_VERIFICADO_CON_EXTREMOS'
        else:raise ValueError(f'{name}: objetivo no identificable de forma única')
        if filenames and target not in filenames:raise ValueError('Nombre de archivo y columna Teléfono discrepan')
        if not any(target in (phone(r[mapping['a']]),phone(r[mapping['b']])) for r in rows):raise ValueError('Objetivo sin respaldo en extremos A/B')
        return headers,mapping,rows,target,{'sheets':{s.title:{'rows':max(0,s.max_row-1),'columns':s.max_column} for s in w},'target_method':method}
    finally:w.close()

def import_files(store,files,options=None):
    options=options or ImportOptions();options.validate()
    existing=store.meta('import_options')
    if existing and json.loads(existing)!=json.loads(dump(options.__dict__)):raise ValueError('Opciones distintas: recalcular explícitamente antes de cambiar política')
    prepared=[];seen={r['sha256'] for r in store.rows('SELECT sha256 FROM imports')}
    skipped=[]
    for name,data in files:
        name=Path(name.replace('\\','/')).name
        sha=hashlib.sha256(data).hexdigest()
        if sha in seen:skipped.append(name);continue
        inspected=inspect_workbook(name,data);prepared.append((name,data,sha,inspected));seen.add(sha)
    if sum(len(p[3][2]) for p in prepared)>500000:
        raise ValueError('El lote V1 excede 500,000 filas; importar en tandas menores')
    if not prepared:return {'imported':0,'skipped':skipped}
    paths=[];counts=Counter()
    try:
        with store.transaction('IMPORTAR_MULTI_CDR',detail={'files':[p[0] for p in prepared],'options':options.__dict__}):
            store.con.execute('INSERT OR REPLACE INTO meta VALUES(?,?)',('import_options',dump(options.__dict__)))
            for name,data,sha,(headers,mapping,rows,target,details) in prepared:
                iid=sha;rel='imports/'+sha+'.xlsx';path=store.path/rel
                if not path.exists():path.write_bytes(data);paths.append(path)
                store.con.execute('INSERT INTO imports VALUES(?,?,?,?,?,?,?,?)',(iid,name,sha,target,now(),rel,dump(options.__dict__),dump(details)))
                ensure_entity(store.con,'TELEFONO',target,target,{'number':target},target=True)
                for n,values in enumerate(rows,2):
                    raw={h:safe_cell(v) for h,v in zip(headers,values)}
                    row={k:(values[i] if i is not None else None) for k,i in mapping.items()}
                    rid=stable(iid,'Datos_Limpios',n)
                    store.con.execute('INSERT INTO raw_records VALUES(?,?,?,?,?)',(rid,iid,'Datos_Limpios',n,dump(raw)))
                    typ=classify(row['kind']);reason='';quality=[]
                    dt=row['datetime']
                    if not text(dt) and row['date'] is not None and row['time'] is not None:
                        dt=str(row['date']).split(' ')[0]+' '+str(row['time'])
                    local,epoch,policy=temporal(dt)
                    a,b=endpoint(row['a']),endpoint(row['b'])
                    aa,bb=a[1],b[1];dur=duration(row['duration']);imei=device(row['imei']);imsi=device(row['imsi'],sim=True)
                    quality.append(policy)
                    if typ=='NO_CLASIFICADO':reason='TIPO_NO_CORRESPONDE_A_TRAFICO: '+text(row['kind'])[:300]
                    elif epoch is None:reason=policy
                    elif target not in (phone(row['a']),phone(row['b'])):reason='OBJETIVO_NO_ES_EXTREMO; no atribuir tráfico a la línea del archivo'
                    elif typ.endswith('ENTRANTE') and phone(row['b'])!=target:reason='DIRECCION_ENTRANTE_CONTRADICTORIA'
                    elif typ.endswith('SALIENTE') and phone(row['a'])!=target:reason='DIRECCION_SALIENTE_CONTRADICTORIA'
                    elif typ.startswith(('VOZ','MENSAJE')) and (not aa or not bb):reason='EXTREMO_NO_INTERPRETABLE'
                    if dur is None:quality.append('DURACION_AUSENTE_O_INVALIDA; no fusionar espejo')
                    if text(row['imei']) and not imei:quality.append('IMEI_NO_UTILIZABLE')
                    if imei and len(imei)==14:quality.append('IMEI_14_DIGITOS_OBSERVADO_SIN_COMPLETAR')
                    if text(row['imsi']) and not imsi:quality.append('IMSI_NO_UTILIZABLE')
                    if (isinstance(row['a'],float) and abs(row['a'])>=10**15) or (isinstance(row['b'],float) and abs(row['b'])>=10**15):reason='PRECISION_NUMERICA_NO_GARANTIZADA'
                    loc=None
                    if not reason:
                        loc,issues=locate(store.con,row);quality.extend(issues)
                        for kind,key in (a,b):
                            if kind and key and not typ.startswith('DATOS'):ensure_entity(store.con,kind,key)
                        if imei:ensure_entity(store.con,'DISPOSITIVO',imei,'IMEI '+imei,{'imei':imei,'digits':len(imei),'validation':'Observado; no completar dígitos'})
                        if imsi:ensure_entity(store.con,'SIM',imsi,'IMSI '+imsi)
                    status='NO_CLASIFICADO' if typ=='NO_CLASIFICADO' else ('REVISAR' if reason else 'VALIDO')
                    # Native duplicate fingerprint includes every delivered field, not merely minute/phone.
                    signature=stable(target,raw)
                    store.con.execute('INSERT INTO observations VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(rid,target,typ,aa,bb,local,epoch,dur,imei or None,imsi or None,loc,status,reason,dump(quality),signature))
                    counts[status]+=1
            from .event_engine import rebuild_events
            from .relation_engine import rebuild_relations
            rebuild_events(store,options);rebuild_relations(store)
            # Existing finding support references events: invalidate and explicitly mark refresh.
            store.con.execute('DELETE FROM findings')
            store.con.execute('INSERT OR REPLACE INTO meta VALUES(?,?)',('correlation_status','PENDIENTE_RECALCULO'))
    except BaseException:
        for p in paths:p.unlink(missing_ok=True)
        raise
    return {'imported':len(prepared),'skipped':skipped,'rows':dict(counts)}
